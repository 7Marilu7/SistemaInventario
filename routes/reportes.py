from flask import render_template, request, send_file
from database import conectar_db
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def generar_excel(titulo, encabezados, filas, nombre_archivo):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    hoja.append(encabezados)

    for celda in hoja[1]:

        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(fill_type="solid", fgColor="252A34")
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila in filas:

        hoja.append([valor.isoformat()if hasattr(valor, "isoformat")else valor for valor in fila])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    hoja.row_dimensions[1].height = 25

    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)
        ancho = max(len(str(celda.value or ""))for celda in columna)
        hoja.column_dimensions[letra].width = min(ancho + 3, 40)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

def registrar_rutas_reportes(app):

    @app.route("/reportes")
    def reportes():

        fecha_inicio = request.args.get("fecha_inicio", "").strip()
        fecha_fin = request.args.get("fecha_fin", "").strip()
        formato_fecha = "%Y-%m-%d"

        if fecha_inicio:
            try:
                datetime.strptime(fecha_inicio, formato_fecha)
            except ValueError:
                fecha_inicio = ""

        if fecha_fin:
            try:
                datetime.strptime(fecha_fin, formato_fecha)
            except ValueError:
                fecha_fin = ""

        if (
            fecha_inicio
            and fecha_fin
            and fecha_inicio > fecha_fin
        ):
            fecha_inicio, fecha_fin = fecha_fin, fecha_inicio

        conexion = conectar_db()

        try:
            
            resumen = conexion.execute("""
                SELECT
                    COUNT(*) AS total_productos,

                    COALESCE(
                        SUM(existencias), 0
                    ) AS total_existencias,

                    COALESCE(
                        SUM(existencias * precio_costo), 0
                    ) AS valor_inventario,

                    COALESCE(
                        SUM(existencias * precio_venta), 0
                    ) AS valor_venta_estimado

                FROM productos

                WHERE activo = 1
            """).fetchone()

            alertas = conexion.execute("""
                SELECT
                    p.id,
                    p.codigo,
                    p.nombre,
                    p.existencias,
                    p.stock_minimo,
                    c.nombre AS categoria

                FROM productos p

                LEFT JOIN categorias c
                    ON c.id = p.categoria_id

                WHERE p.activo = 1
                  AND p.existencias <= p.stock_minimo

                ORDER BY
                    p.existencias ASC,
                    p.nombre COLLATE NOCASE
            """).fetchall()

            inventario = conexion.execute("""
                SELECT
                    p.codigo,
                    p.nombre,

                    c.nombre AS categoria,

                    p.existencias,
                    p.stock_minimo,
                    p.precio_costo,
                    p.precio_venta,

                    p.existencias * p.precio_costo
                        AS valor_costo,

                    p.existencias * p.precio_venta
                        AS valor_venta

                FROM productos p

                LEFT JOIN categorias c
                    ON c.id = p.categoria_id

                WHERE p.activo = 1

                ORDER BY p.nombre COLLATE NOCASE
            """).fetchall()

            consulta_movimientos = """
                SELECT
                    m.id,
                    m.tipo,
                    m.cantidad,
                    m.existencias_anteriores,
                    m.existencias_nuevas,
                    m.motivo,
                    m.notas,
                    m.fecha,
                    p.codigo,
                    p.nombre AS producto

                FROM movimientos m

                INNER JOIN productos p
                    ON p.id = m.producto_id

                WHERE 1 = 1
            """

            parametros = []

            if fecha_inicio:
                consulta_movimientos += """
                    AND date(m.fecha) >= date(?)
                """
                parametros.append(fecha_inicio)

            if fecha_fin:
                consulta_movimientos += """
                    AND date(m.fecha) <= date(?)
                """
                parametros.append(fecha_fin)

            consulta_movimientos += """
                ORDER BY m.fecha DESC, m.id DESC
            """

            movimientos = conexion.execute(
                consulta_movimientos,
                parametros
            ).fetchall()

            consulta_totales = """
                SELECT

                    COALESCE(
                        SUM(
                            CASE
                                WHEN tipo = 'entrada'
                                THEN cantidad
                                ELSE 0
                            END
                        ),
                        0
                    ) AS total_entradas,

                    COALESCE(
                        SUM(
                            CASE
                                WHEN tipo = 'salida'
                                THEN cantidad
                                ELSE 0
                            END
                        ),
                        0
                    ) AS total_salidas,

                    COUNT(*) AS total_movimientos

                FROM movimientos

                WHERE 1 = 1
            """

            parametros_totales = []

            if fecha_inicio:
                consulta_totales += """
                    AND date(fecha) >= date(?)
                """

                parametros_totales.append(fecha_inicio)

            if fecha_fin:
                consulta_totales += """
                    AND date(fecha) <= date(?)
                """

                parametros_totales.append(fecha_fin)

            totales_movimientos = conexion.execute(
                consulta_totales,
                parametros_totales
            ).fetchone()

            return render_template(
                "reportes.html",
                resumen=resumen,
                alertas=alertas,
                inventario=inventario,
                movimientos=movimientos,
                totales_movimientos=totales_movimientos,
                fecha_inicio=fecha_inicio,
                fecha_fin=fecha_fin
            )

        finally:
            conexion.close()

    @app.route("/reportes/exportar/inventario")
    def exportar_inventario_excel():

        conexion = conectar_db()

        try:

            registros = conexion.execute("""
                SELECT
                    p.codigo,
                    p.nombre,

                    COALESCE(
                        c.nombre,
                        'Sin categoría'
                    ) AS categoria,

                    p.existencias,
                    p.stock_minimo,
                    p.precio_costo,
                    p.precio_venta,

                    p.existencias * p.precio_costo
                        AS valor_costo,

                    p.existencias * p.precio_venta
                        AS valor_venta

                FROM productos p

                LEFT JOIN categorias c
                    ON c.id = p.categoria_id

                WHERE p.activo = 1

                ORDER BY p.nombre COLLATE NOCASE
            """).fetchall()

            encabezados = [
                "Código",
                "Producto",
                "Categoría",
                "Existencias",
                "Stock mínimo",
                "Precio de costo",
                "Precio de venta",
                "Valor al costo",
                "Valor de venta"
            ]

            filas = [
                tuple(registro)
                for registro in registros
            ]

            return generar_excel(
                "Inventario",
                encabezados,
                filas,
                "inventario.xlsx"
            )

        finally:
            conexion.close()

    @app.route("/reportes/exportar/movimientos")
    def exportar_movimientos_excel():

        conexion = conectar_db()

        try:

            registros = conexion.execute("""
                SELECT
                    m.fecha,
                    m.tipo,
                    p.codigo,
                    p.nombre AS producto,
                    m.cantidad,
                    m.existencias_anteriores,
                    m.existencias_nuevas,
                    m.motivo,
                    m.notas

                FROM movimientos m

                INNER JOIN productos p
                    ON p.id = m.producto_id

                ORDER BY m.fecha DESC, m.id DESC
            """).fetchall()

            encabezados = [
                "Fecha",
                "Tipo",
                "Código",
                "Producto",
                "Cantidad",
                "Existencias anteriores",
                "Existencias nuevas",
                "Motivo",
                "Notas"
            ]

            filas = [
                tuple(registro)
                for registro in registros
            ]

            return generar_excel(
                "Movimientos",
                encabezados,
                filas,
                "movimientos.xlsx"
            )

        finally:
            conexion.close()

    @app.route("/reportes/exportar/movimientos-filtrados")
    def exportar_movimientos_filtrados_excel():

        fecha_inicio = request.args.get("fecha_inicio", "").strip()
        fecha_fin = request.args.get("fecha_fin", "").strip()
        formato_fecha = "%Y-%m-%d"

        if fecha_inicio:
            try:
                datetime.strptime(
                    fecha_inicio,
                    formato_fecha
                )
            except ValueError:
                fecha_inicio = ""

        if fecha_fin:
            try:
                datetime.strptime(
                    fecha_fin,
                    formato_fecha
                )
            except ValueError:
                fecha_fin = ""

        if (
            fecha_inicio
            and fecha_fin
            and fecha_inicio > fecha_fin
        ):
            fecha_inicio, fecha_fin = fecha_fin, fecha_inicio

        consulta = """
            SELECT
                m.fecha,
                m.tipo,
                p.codigo,
                p.nombre AS producto,
                m.cantidad,
                m.existencias_anteriores,
                m.existencias_nuevas,
                m.motivo,
                m.notas

            FROM movimientos m

            INNER JOIN productos p
                ON p.id = m.producto_id

            WHERE 1 = 1
        """

        parametros = []

        if fecha_inicio:
            consulta += """
                AND date(m.fecha) >= date(?)
            """

            parametros.append(fecha_inicio)

        if fecha_fin:
            consulta += """
                AND date(m.fecha) <= date(?)
            """

            parametros.append(fecha_fin)

        consulta += """
            ORDER BY m.fecha DESC, m.id DESC
        """

        conexion = conectar_db()

        try:

            registros = conexion.execute(
                consulta,
                parametros
            ).fetchall()

            encabezados = [
                "Fecha",
                "Tipo",
                "Código",
                "Producto",
                "Cantidad",
                "Existencias anteriores",
                "Existencias nuevas",
                "Motivo",
                "Notas"
            ]

            filas = [
                tuple(registro)
                for registro in registros
            ]

            return generar_excel(
                "Movimientos",
                encabezados,
                filas,
                "movimientos_filtrados.xlsx"
            )

        finally:
            conexion.close()